# /// script
# dependencies = ["python-docx"]
# ///
# Minimal Markdown -> .docx for our report subset: # ## ### headings, | tables |, - bullets,
# 1. numbered, **bold**, `code`, --- rule, paragraphs. Run: uv run scripts/md_to_docx.py in.md out.docx
import sys, re
from docx import Document
from docx.shared import Pt, RGBColor

src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

def add_runs(par, text):
    # split on **bold** and `code`, emit runs
    for seg in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = par.add_run(seg[2:-2]); r.bold = True
        elif seg.startswith("`") and seg.endswith("`"):
            r = par.add_run(seg[1:-1]); r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            par.add_run(seg)

i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1; continue
    if ln.startswith("### "):
        doc.add_heading(ln[4:], level=3)
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=2)
    elif ln.startswith("# "):
        doc.add_heading(ln[2:], level=1)
    elif re.match(r"^-{3,}$", ln):
        pass  # horizontal rule -> skip (visual only)
    elif ln.lstrip().startswith("- "):
        add_runs(doc.add_paragraph(style="List Bullet"), ln.lstrip()[2:])
    elif re.match(r"^\d+\.\s", ln.lstrip()):
        add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", ln.lstrip()))
    elif ln.lstrip().startswith("|"):
        # gather the table block
        block = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            block.append(lines[i].strip()); i += 1
        rows = [[c.strip() for c in r.strip("|").split("|")] for r in block]
        rows = [r for r in rows if not all(re.match(r"^:?-{2,}:?$", c or "-") for c in r)]  # drop --- separator
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    if ci < len(t.rows[ri].cells):
                        cellp = t.rows[ri].cells[ci].paragraphs[0]; add_runs(cellp, cell)
                        if ri == 0:
                            for run in cellp.runs: run.bold = True
        continue
    else:
        add_runs(doc.add_paragraph(), ln)
    i += 1

doc.save(out)
print("wrote", out)
